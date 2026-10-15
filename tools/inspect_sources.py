# -*- coding: utf-8 -*-
"""
Скрипт разведки исходных данных (раздел 3 ТЗ).
Показывает по каждому файлу: кодировку, BOM, разделитель, заголовок и первые строки.
Данные НЕ изменяет — только читает.
"""
import os
import sys
import codecs
import csv
import contextlib

ROOT = sys.argv[1] if len(sys.argv) > 1 else \
    r'C:\Users\user\Downloads\AyuGram Desktop\servicedesk_data_unpacked\servicedesk_data'


def detect_encoding(raw: bytes):
    if raw.startswith(codecs.BOM_UTF8):
        return 'utf-8-sig', 'UTF-8 BOM'
    if raw.startswith(codecs.BOM_UTF16_LE):
        return 'utf-16', 'UTF-16 LE BOM'
    if raw.startswith(codecs.BOM_UTF16_BE):
        return 'utf-16', 'UTF-16 BE BOM'
    for enc in ('utf-8', 'cp1251'):
        try:
            raw.decode(enc)
            return enc, 'нет BOM'
        except UnicodeDecodeError:
            continue
    return 'cp1251', 'не удалось определить'


def guess_delimiter(line: str):
    # разделитель — тот, которого больше всего в первой строке
    counts = {d: line.count(d) for d in (';', ',', '\t', '|')}
    best = max(counts, key=counts.get)
    return best, counts


def inspect_csv(path: str):
    with open(path, 'rb') as f:
        raw = f.read()
    enc, bom = detect_encoding(raw)
    text = raw.decode(enc, errors='replace')
    lines = text.splitlines()
    header = lines[0] if lines else ''
    delim, counts = guess_delimiter(header)
    print(f'  Кодировка : {enc}  ({bom})')
    print(f'  Разделитель: "{delim}"  (в заголовке: {counts})')
    print(f'  Физических строк: {len(lines)}')
    print(f'  Заголовок : {header}')
    print('  Первые 4 строки данных:')
    for ln in lines[1:5]:
        print('    | ' + ln[:300])
    # число записей csv и признак многострочных полей
    try:
        with open(path, 'r', encoding=enc, newline='') as fh:
            reader = csv.reader(fh, delimiter=delim)
            rows = list(reader)
        print(f'  Логических записей (csv): {len(rows)}')
        if len(rows) != len(lines):
            print(f'  !!! Записей csv ({len(rows)}) != физических строк ({len(lines)}) '
                  f'— есть многострочные поля/переносы')
        widths = {len(r) for r in rows[:2000]}
        print(f'  Колонок в записях (набор): {sorted(widths)}')
    except Exception as e:
        print(f'  Не удалось распарсить как CSV: {e}')


def inspect_text(path: str, maxlines: int = 20):
    with open(path, 'rb') as f:
        raw = f.read()
    enc, bom = detect_encoding(raw)
    text = raw.decode(enc, errors='replace')
    lines = text.splitlines()
    print(f'  Кодировка : {enc}  ({bom})')
    print(f'  Строк     : {len(lines)}')
    print(f'  Первые {maxlines} строк:')
    for ln in lines[:maxlines]:
        print('    | ' + ln[:300])


def inspect_xlsx(path: str):
    try:
        import openpyxl
    except ImportError:
        print('  openpyxl не установлен — установлю позже для чтения XLSX')
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f'  Лист "{ws.title}": {ws.max_row} строк x {ws.max_column} столбцов')
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 6:
                break
            print('    | ' + ' ; '.join('' if c is None else str(c) for c in row)[:300])
    wb.close()


def main():
    print('=' * 78)
    print('РАЗВЕДКА ИСХОДНЫХ ДАННЫХ: ' + ROOT)
    print('=' * 78)

    groups = {
        'ref': ['categories.csv', 'priorities_sla.csv', 'production_calendar.csv',
                'engineer_skills.csv', 'duty_schedule.csv'],
        'hr': ['departments.csv', 'employees.csv'],
        'cmdb': ['asset_movements.csv'],
        'legacy': ['tickets.csv', 'ticket_status_history.csv'],
    }

    for folder, names in groups.items():
        for name in names:
            path = os.path.join(ROOT, folder, name)
            print('-' * 78)
            print(f'[CSV] {folder}/{name}')
            print('-' * 78)
            if os.path.exists(path):
                inspect_csv(path)
            else:
                print('  ФАЙЛ НЕ НАЙДЕН')
            print()

    # XLSX
    xlsx = os.path.join(ROOT, 'cmdb', 'assets_1c_export.xlsx')
    print('-' * 78)
    print('[XLSX] cmdb/assets_1c_export.xlsx')
    print('-' * 78)
    if os.path.exists(xlsx):
        inspect_xlsx(xlsx)
    print()

    # LDIF
    ldif = os.path.join(ROOT, 'ldap', 'directory.ldif')
    print('-' * 78)
    print('[LDIF] ldap/directory.ldif')
    print('-' * 78)
    if os.path.exists(ldif):
        inspect_text(ldif, 25)
    print()

    # Один пример письма
    eml = os.path.join(ROOT, 'mail', '025.eml')
    print('-' * 78)
    print('[EML] mail/025.eml (пример)')
    print('-' * 78)
    if os.path.exists(eml):
        inspect_text(eml, 30)
    print()


if __name__ == '__main__':
    report_path = os.path.abspath(os.path.join(
        os.path.dirname(__file__), '..', 'docs', 'source_inspection.txt'))
    os.makedirs(os.path.dirname(report_path), exist_ok=True)
    with open(report_path, 'w', encoding='utf-8') as f, contextlib.redirect_stdout(f):
        main()
    print('Отчёт сохранён: ' + report_path)
